"""Zwei Folien zur Druck-Temperatur-Kopplung.  Aufruf aus figures/:  python3 make_slides_pressure.py"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

BLUE = RGBColor(0x2a, 0x78, 0xd6); GREY = RGBColor(0x52, 0x51, 0x4e)
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)

def slide(title):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.8))
    p = tb.text_frame.paragraphs[0]; p.text = title
    p.font.size = Pt(30); p.font.bold = True; p.font.color.rgb = GREY
    return s

def bullets(s, left, top, width, height, lines, size=15):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    for i, (txt, lvl, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt; p.level = lvl
        p.font.size = Pt(size); p.font.bold = bold
        p.font.color.rgb = BLUE if bold else GREY
        p.space_after = Pt(6)
    return tb

# ---------------------------------------------------------------- Folie 1
s1 = slide("Firn temperature follows barometric pressure")
s1.shapes.add_picture("fig3c_rohdaten.png", Inches(0.4), Inches(1.15), width=Inches(7.6))
s1.shapes.add_picture("fig3a_kopplung_tiefe.png", Inches(8.1), Inches(1.15), width=Inches(4.9))
bullets(s1, 0.5, 5.9, 12.4, 1.4, [
 ("Raw data (left): firn temperature vs. air pressure, synoptic band. Two sites, "
  "one line per depth — the coupling is in the data, not in the statistics.", 0, False),
 ("Depth dependence (right): measured against the parameter-free prediction "
  "b = φ/(ρ c), with no free parameter. Factor 0.95 ± 0.02 over 10 nodes, 7–30 m.", 0, False),
 ("Kohnen B50 (6 months, 41 hPa) · GRIP double chain (2 months, 21 hPa) · "
  "Kohnen B46 (2 weeks, open symbols: shape only)", 0, False)])

# ---------------------------------------------------------------- Folie 2
s2 = slide("A route to in-situ open porosity and close-off")
bullets(s2, 0.5, 1.2, 6.2, 5.6, [
 ("Principle", 0, True),
 ("A pressure rise pushes air into the connected pore space; the compression heat "
  "is absorbed by the ice matrix:", 0, False),
 ("b = ΔT/Δp = φ_open / (ρ · c_ice)", 0, True),
 ("Only pore space connected to the atmosphere responds — the functional quantity "
  "firn-air models need, not the geometric one from a core.", 0, False),
 ("", 0, False),
 ("Two read-outs", 0, True),
 ("Backfilled hole: in-phase response → φ_open(z). With core density "
  "(φ_total = 1 − ρ/ρ_ice) the difference is the closed porosity.", 0, False),
 ("Open borehole: response to dp/dt drops 4.5× between 75 and 92 m at B40 — "
  "the wall turns impermeable. The hole marks its own close-off.", 0, False)])
bullets(s2, 6.9, 1.2, 6.0, 5.6, [
 ("Expected magnitudes (Kohnen)", 0, True),
 ("0.09 mK/hPa at 1 m → 0.03 at 30 m → 0.006 at close-off (84 m); "
  "below: thermoelastic floor of ice, 0.002 mK/hPa", 0, False),
 ("", 0, False),
 ("Precision", 0, True),
 ("0.1 mK per sample, σ_p = 5 hPa, 2 samples/day → φ_open to 0.6 % at 10 m, "
  "5 % at close-off after 3 years", 0, False),
 ("Lab pycnometry on cores: 2 % at low density, up to 20 % near close-off, "
  "plus cut-bubble and relaxation corrections", 0, False),
 ("", 0, False),
 ("What it needs", 0, True),
 ("Thermistor chain in a backfilled hole through the lock-in zone · reference "
  "barometer within ~100 km, time-matched · ≥2 samples/day · 1–3 years", 0, False),
 ("", 0, False),
 ("Limits", 0, True),
 ("Vertical resolution = node spacing (metres) — complementary to μCT, not a "
  "replacement. Absolute accuracy set by c_ice (~1 %).", 0, False)])

prs.save("slides_druck_temperatur.pptx")
print("slides_druck_temperatur.pptx geschrieben")
